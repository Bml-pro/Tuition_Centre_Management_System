
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from students.models import Student
from .forms import PaymentForm
from .models import Payment
from payments.services import get_monthly_tuition


def get_report_data(request):
    """Build the monthly report data used by the page and exports."""

    today = timezone.localdate()

    try:
        month = int(request.GET.get("month", today.month))
    except (ValueError, TypeError):
        month = today.month

    try:
        year = int(request.GET.get("year", today.year))
    except (ValueError, TypeError):
        year = today.year

    if month < 1 or month > 12:
        month = today.month

    if year < 1900 or year > 9999:
        year = today.year

    monthly_fee = Decimal(str(get_monthly_tuition() or 0))

    students = Student.objects.filter(
        status="ACTIVE"
    ).order_by("full_name")

    report = []
    total_expected = Decimal("0.00")
    total_paid = Decimal("0.00")

    for student in students:
        paid = (
            Payment.objects.filter(
                student=student,
                payment_month=month,
                payment_year=year,
            ).aggregate(total=Sum("amount"))["total"]
            or Decimal("0.00")
        )

        balance = max(monthly_fee - paid, Decimal("0.00"))

        if monthly_fee <= 0:
            status = "FEE NOT SET"
        elif paid >= monthly_fee:
            status = "PAID"
        elif paid > 0:
            status = "PARTIALLY PAID"
        else:
            status = "UNPAID"

        report.append({
            "student": student,
            "monthly_fee": monthly_fee,
            "paid": paid,
            "balance": balance,
            "status": status,
        })

        total_expected += monthly_fee
        total_paid += paid

    total_balance = max(
        total_expected - total_paid,
        Decimal("0.00"),
    )

    month_name = dict(Payment.MONTH_CHOICES).get(month, "Unknown")

    return {
        "report": report,
        "selected_month": month,
        "selected_year": year,
        "month_name": month_name,
        "monthly_fee": monthly_fee,
        "total_expected": total_expected,
        "total_paid": total_paid,
        "total_balance": total_balance,
        "months": Payment.MONTH_CHOICES,
    }


@login_required
def payment_list(request):
    payments = Payment.objects.select_related(
        "student",
        "recorded_by",
    ).order_by(
        "-payment_year",
        "-payment_month",
        "-payment_date",
        "-created_at",
    )

    search = request.GET.get("search", "").strip()
    selected_month = request.GET.get("month", "").strip()
    selected_year = request.GET.get("year", "").strip()
    selected_student = request.GET.get("student", "").strip()

    if search:
        payments = payments.filter(
            Q(student__full_name__icontains=search)
            | Q(student__student_id__icontains=search)
        )

    if selected_month:
        try:
            month_number = int(selected_month)
            if 1 <= month_number <= 12:
                payments = payments.filter(
                    payment_month=month_number
                )
            else:
                selected_month = ""
        except ValueError:
            selected_month = ""

    if selected_year:
        try:
            payments = payments.filter(
                payment_year=int(selected_year)
            )
        except ValueError:
            selected_year = ""

    if selected_student:
        try:
            payments = payments.filter(
                student_id=int(selected_student)
            )
        except ValueError:
            selected_student = ""

    total_amount = (
        payments.aggregate(total=Sum("amount"))["total"]
        or Decimal("0.00")
    )
    payment_count = payments.count()

    students = Student.objects.order_by("full_name")

    years = (
        Payment.objects.values_list("payment_year", flat=True)
        .distinct()
        .order_by("-payment_year")
    )

    paginator = Paginator(payments, 15)
    page_obj = paginator.get_page(request.GET.get("page"))

    context = {
        "payments": page_obj,
        "page_obj": page_obj,
        "students": students,
        "years": years,
        "months": Payment.MONTH_CHOICES,
        "search": search,
        "selected_month": selected_month,
        "selected_year": selected_year,
        "selected_student": selected_student,
        "total_amount": total_amount,
        "payment_count": payment_count,
        "monthly_fee": get_monthly_tuition(),
    }

    return render(request, "payments/payment_list.html", context)


@login_required
def monthly_payment_report(request):
    context = get_report_data(request)
    return render(
        request,
        "payments/monthly_report.html",
        context,
    )


@login_required
def export_monthly_report_excel(request):
    data = get_report_data(request)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Monthly Tuition Report"

    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = (
        f"MONTHLY TUITION REPORT - "
        f"{data['month_name']} {data['selected_year']}"
    )
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )
    worksheet["A1"].fill = PatternFill(
        "solid",
        fgColor="1D4ED8",
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    worksheet.append([])
    worksheet.append([
        "Student ID",
        "Student Name",
        "Monthly Fee (TSh)",
        "Amount Paid (TSh)",
        "Balance (TSh)",
        "Status",
    ])

    header_row = 3

    for cell in worksheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center")

    for item in data["report"]:
        worksheet.append([
            item["student"].student_id,
            item["student"].full_name,
            float(item["monthly_fee"]),
            float(item["paid"]),
            float(item["balance"]),
            item["status"],
        ])

    summary_row = worksheet.max_row + 2
    worksheet.cell(summary_row, 2, "TOTAL EXPECTED")
    worksheet.cell(
        summary_row, 3, float(data["total_expected"])
    )
    worksheet.cell(summary_row + 1, 2, "TOTAL COLLECTED")
    worksheet.cell(
        summary_row + 1, 3, float(data["total_paid"])
    )
    worksheet.cell(summary_row + 2, 2, "TOTAL BALANCE")
    worksheet.cell(
        summary_row + 2, 3, float(data["total_balance"])
    )

    for row in range(summary_row, summary_row + 3):
        worksheet.cell(row, 2).font = Font(bold=True)
        worksheet.cell(row, 3).number_format = '#,##0.00'

    for row in range(4, 4 + len(data["report"])):
        for column in (3, 4, 5):
            worksheet.cell(row, column).number_format = '#,##0.00'

    column_widths = {
        "A": 18,
        "B": 30,
        "C": 20,
        "D": 20,
        "E": 18,
        "F": 22,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = (
        f"A3:F{3 + len(data['report'])}"
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        f"tuition_report_{data['selected_year']}_"
        f"{data['selected_month']:02d}.xlsx"
    )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )
    return response


@login_required
def export_monthly_report_pdf(request):
    data = get_report_data(request)

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1D4ED8"),
        fontSize=17,
        spaceAfter=8,
    )

    story = [
        Paragraph(
            "WAMBURA TUITION SYSTEM",
            title_style,
        ),
        Paragraph(
            f"Monthly Tuition Report — "
            f"{data['month_name']} {data['selected_year']}",
            styles["Heading2"],
        ),
        Spacer(1, 5 * mm),
        Paragraph(
            f"Monthly tuition fee: "
            f"TSh {data['monthly_fee']:,.2f}",
            styles["Normal"],
        ),
        Spacer(1, 5 * mm),
    ]

    table_data = [[
        "Student ID",
        "Student Name",
        "Monthly Fee (TSh)",
        "Paid (TSh)",
        "Balance (TSh)",
        "Status",
    ]]

    for item in data["report"]:
        table_data.append([
            item["student"].student_id,
            Paragraph(
                item["student"].full_name,
                styles["Normal"],
            ),
            f"{item['monthly_fee']:,.2f}",
            f"{item['paid']:,.2f}",
            f"{item['balance']:,.2f}",
            item["status"],
        ])

    table = Table(
        table_data,
        colWidths=[
            28 * mm,
            65 * mm,
            36 * mm,
            34 * mm,
            34 * mm,
            35 * mm,
        ],
        repeatRows=1,
    )

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor("#F1F5F9"),
        ]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    story.append(table)
    story.append(Spacer(1, 8 * mm))

    summary_data = [
        ["Total Expected", "Total Collected", "Total Balance"],
        [
            f"TSh {data['total_expected']:,.2f}",
            f"TSh {data['total_paid']:,.2f}",
            f"TSh {data['total_balance']:,.2f}",
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[65 * mm, 65 * mm, 65 * mm],
    )
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))

    story.append(summary_table)

    document.build(story)
    output.seek(0)

    filename = (
        f"tuition_report_{data['selected_year']}_"
        f"{data['selected_month']:02d}.pdf"
    )

    response = HttpResponse(
        output.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )
    return response


@login_required
def payment_create(request):
    student_id = request.GET.get("student")
    selected_student = None
    initial_data = {}

    if student_id:
        try:
            selected_student = Student.objects.get(pk=student_id)
            initial_data["student"] = selected_student
        except (Student.DoesNotExist, ValueError):
            pass

    if request.method == "POST":
        form = PaymentForm(request.POST)

        if form.is_valid():
            payment = form.save(commit=False)
            payment.recorded_by = request.user
            payment.save()

            messages.success(
                request,
                f"Payment of TSh {payment.amount:,.2f} "
                f"for {payment.student.full_name} was recorded successfully.",
            )

            return redirect("payment_detail", pk=payment.pk)
    else:
        form = PaymentForm(initial=initial_data)

    return render(
        request,
        "payments/payment_form.html",
        {
            "form": form,
            "page_title": "Record Payment",
            "selected_student": selected_student,
        },
    )


@login_required
def payment_detail(request, pk):
    payment = get_object_or_404(
        Payment.objects.select_related("student", "recorded_by"),
        pk=pk,
    )

    return render(
        request,
        "payments/payment_detail.html",
        {"payment": payment},
    )
